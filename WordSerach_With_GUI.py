"""Search Word, Excel, PowerPoint, and PDF files by keyword.

Dependencies:
    pip install python-docx openpyxl python-pptx python-Levenshtein PyMuPDF

The GUI's minimum keyword score counts DISTINCT search terms found anywhere
in one document, not the number of occurrences of a term.
"""

import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, Toplevel

from docx import Document
import openpyxl
from pptx import Presentation
import Levenshtein
import fitz  # PyMuPDF for PDF support


# ---------------- Matching Logic ------------------
def is_match(word, keyword, max_mismatches, use_substring_match, use_blurry):
    word_lower = word.lower()
    keyword_lower = keyword.lower()

    if use_substring_match and keyword_lower in word_lower:
        return True

    if use_blurry:
        return Levenshtein.distance(word_lower, keyword_lower) <= max_mismatches

    return word_lower == keyword_lower


# ---------------- Text Extraction ------------------
def extract_text_from_docx(path):
    try:
        doc = Document(path)
        return "\n".join(para.text for para in doc.paragraphs)
    except Exception as exc:
        print(f"Error reading DOCX {path}: {exc}")
        return ""


def extract_text_from_xlsx(path):
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            text = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            text.append(str(cell.value))
            return " ".join(text)
        finally:
            wb.close()
    except Exception as exc:
        print(f"Error reading XLSX {path}: {exc}")
        return ""


def extract_text_from_pptx(path):
    try:
        prs = Presentation(path)
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text.append(shape.text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            text.append(cell.text)
        return " ".join(text)
    except Exception as exc:
        print(f"Error reading PPTX {path}: {exc}")
        return ""


def extract_text_from_pdf(path):
    try:
        with fitz.open(path) as doc:
            return " ".join(page.get_text() for page in doc)
    except Exception as exc:
        print(f"Error reading PDF {path}: {exc}")
        return ""


def extract_text(path):
    extension = os.path.splitext(path)[1].lower()
    if extension == ".docx":
        return extract_text_from_docx(path)
    elif extension == ".xlsx":
        return extract_text_from_xlsx(path)
    elif extension == ".pptx":
        return extract_text_from_pptx(path)
    elif extension == ".pdf":
        return extract_text_from_pdf(path)
    return ""


def find_matches(text, keywords, max_mismatches, allow_partial, use_blurry,
                 context_window=3):
    """Return unique keywords found, matching contexts, and hit count.

    Each matched word gets one context snippet. Exact matches take priority
    over substring matches, which take priority over fuzzy matches. This keeps
    a single word from satisfying multiple SIMILAR fuzzy keywords by accident.
    """
    words = text.split()
    matched_keywords = set()
    snippets = []
    total_hits = 0

    for i, word in enumerate(words):
        lowered = word.lower()
        exact_matches = [kw for kw in keywords if kw.lower() == lowered]
        if exact_matches:
            matching_terms = exact_matches
        else:
            substring_matches = (
                [kw for kw in keywords if kw.lower() in lowered]
                if allow_partial else []
            )
            if substring_matches:
                matching_terms = substring_matches
            elif use_blurry:
                distances = [
                    (Levenshtein.distance(lowered, kw.lower()), kw)
                    for kw in keywords
                ]
                best_distance, best_keyword = min(distances)
                matching_terms = ([best_keyword] if best_distance <= max_mismatches
                                  else [])
            else:
                matching_terms = []

        if not matching_terms:
            continue

        matched_keywords.update(term.lower() for term in matching_terms)
        total_hits += 1
        start = max(0, i - context_window)
        end = min(len(words), i + context_window + 1)
        context = words[start:end].copy()
        context[i - start] = f"[{context[i - start]}]"
        snippets.append(" ".join(context))

    return matched_keywords, snippets, total_hits


# ---------------- GUI App ------------------
class SearchApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Document Keyword Search")
        self.root.geometry("1200x700")

        # Store results independently from displayed labels, which include scores.
        self.results = {}
        self.result_paths = []
        self.keyword_count = 0
        self.keyword_display_names = {}

        # Frames
        self.left_frame = tk.Frame(root, width=230)
        self.left_frame.pack(side="left", fill="y", padx=5, pady=5)

        self.middle_frame = tk.Frame(root, width=380)
        self.middle_frame.pack(side="left", fill="both", expand=True, padx=5, pady=5)

        self.right_frame = tk.Frame(root, width=500)
        self.right_frame.pack(side="left", fill="both", expand=True, padx=5, pady=5)

        # Left Panel - Directory + Keywords
        tk.Label(self.left_frame, text="Directory:",
                 font=("TkDefaultFont", 9, "bold")).pack(anchor="w", padx=5)

        self.dir_entry = tk.Entry(self.left_frame, width=30)
        self.dir_entry.pack(padx=5, pady=2)

        tk.Button(self.left_frame, text="Browse",
                  command=self.browse_directory).pack(pady=2)

        # File type filters
        filetypes_frame = tk.Frame(self.left_frame)
        filetypes_frame.pack(pady=2)
        self.include_docx = tk.BooleanVar(value=True)
        self.include_xlsx = tk.BooleanVar(value=True)
        self.include_pptx = tk.BooleanVar(value=True)
        self.include_pdf = tk.BooleanVar(value=True)
        tk.Checkbutton(filetypes_frame, text="Word (.docx)",
                       variable=self.include_docx).pack(anchor="w")
        tk.Checkbutton(filetypes_frame, text="Excel (.xlsx)",
                       variable=self.include_xlsx).pack(anchor="w")
        tk.Checkbutton(filetypes_frame, text="PowerPoint (.pptx)",
                       variable=self.include_pptx).pack(anchor="w")
        tk.Checkbutton(filetypes_frame, text="PDF (.pdf)",
                       variable=self.include_pdf).pack(anchor="w")

        # Search options
        tk.Label(self.left_frame, text="Search Options:",
                 font=("TkDefaultFont", 9, "bold")).pack(
            anchor="w", padx=5, pady=5)

        search_opts_frame = tk.Frame(self.left_frame)
        search_opts_frame.pack(pady=5, anchor="w")

        blurry_row = tk.Frame(search_opts_frame)
        blurry_row.pack(anchor="w", padx=5)
        self.use_blurry = tk.BooleanVar(value=True)
        tk.Checkbutton(blurry_row, text="Blurry search",
                       variable=self.use_blurry).pack(side="left")

        self.max_mismatch_entry = tk.Entry(blurry_row, width=5)
        self.max_mismatch_entry.insert(0, "1")
        self.max_mismatch_entry.pack(side="left", padx=(5, 0))

        self.use_substring_match = tk.BooleanVar(value=True)
        tk.Checkbutton(search_opts_frame, text="Allow partial match",
                       variable=self.use_substring_match).pack(
            anchor="w", padx=5)

        # NEW: minimum number of distinct keywords required for a file.
        minimum_row = tk.Frame(search_opts_frame)
        minimum_row.pack(anchor="w", padx=5, pady=(9, 0))
        tk.Label(minimum_row, text="Minimum keyword score:").pack(side="left")
        self.min_keywords_entry = ttk.Spinbox(
            minimum_row, from_=1, to=999, width=5
        )
        self.min_keywords_entry.set("1")
        self.min_keywords_entry.pack(side="left", padx=(5, 0))

        tk.Label(
            search_opts_frame,
            text="Counts different keywords, not repeats",
            font=("TkDefaultFont", 8), fg="gray35"
        ).pack(anchor="w", padx=5)

        tk.Button(self.left_frame, text="Search",
                  command=self.search_documents).pack(pady=10)

        tk.Label(self.left_frame, text="Keywords:",
                 font=("TkDefaultFont", 9, "bold")).pack(
            anchor="w", padx=5)

        keyword_control_frame = tk.Frame(self.left_frame)
        keyword_control_frame.pack(padx=5, pady=2, anchor="center")
        tk.Button(keyword_control_frame, text="+", width=5,
                  command=self.add_keyword_entry).pack(side="left", padx=2)
        tk.Button(keyword_control_frame, text="-", width=5,
                  command=self.remove_keyword_entry).pack(side="left", padx=2)

        self.keyword_entries = []
        self.add_keyword_entry()

        # Middle Panel - Results, sorted by keyword score
        tk.Label(self.middle_frame, text="Matching documents (highest score first)",
                 font=("TkDefaultFont", 9, "bold")).pack(anchor="w")
        self.result_listbox = tk.Listbox(self.middle_frame, width=55)
        self.result_listbox.pack(side="left", fill="both", expand=True)
        self.result_listbox.bind("<<ListboxSelect>>", self.on_file_select)
        result_scrollbar = ttk.Scrollbar(self.middle_frame, orient="vertical",
                                         command=self.result_listbox.yview)
        result_scrollbar.pack(side="right", fill="y")
        self.result_listbox.configure(yscrollcommand=result_scrollbar.set)

        # Right Panel - Matching contexts
        tk.Label(self.right_frame, text="Match Context",
                 font=("TkDefaultFont", 9, "bold")).pack(anchor="w")
        self.context_text = tk.Text(self.right_frame, wrap="word")
        self.context_text.pack(side="left", fill="both", expand=True)
        context_scrollbar = ttk.Scrollbar(self.right_frame, orient="vertical",
                                          command=self.context_text.yview)
        context_scrollbar.pack(side="right", fill="y")
        self.context_text.configure(yscrollcommand=context_scrollbar.set)

    def add_keyword_entry(self):
        entry = tk.Entry(self.left_frame, width=28)
        entry.pack(padx=5, pady=2)
        self.keyword_entries.append(entry)

    def remove_keyword_entry(self):
        if len(self.keyword_entries) > 1:
            self.keyword_entries.pop().destroy()

    def get_keywords(self):
        # Don't score a duplicate keyword twice (case-insensitive).
        keywords = []
        seen = set()
        for entry in self.keyword_entries:
            keyword = entry.get().strip()
            if keyword and keyword.lower() not in seen:
                seen.add(keyword.lower())
                keywords.append(keyword)
        return keywords

    def browse_directory(self):
        directory = filedialog.askdirectory(title="Select Directory")
        if directory:
            self.dir_entry.delete(0, tk.END)
            self.dir_entry.insert(0, directory)

    def search_documents(self):
        directory = self.dir_entry.get().strip()
        if not directory or not os.path.isdir(directory):
            messagebox.showwarning("Invalid Directory",
                                   "Please select a valid directory.")
            return

        keywords = self.get_keywords()
        if not keywords:
            messagebox.showwarning("No Keywords",
                                   "Please enter at least one keyword.")
            return

        try:
            max_mismatches = int(self.max_mismatch_entry.get())
            if max_mismatches < 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Invalid Blurry Search Setting",
                                   "Allowed mismatches must be 0 or a positive integer.")
            return

        # NEW: validate the minimum keyword score chosen in the GUI.
        try:
            min_keywords_required = int(self.min_keywords_entry.get())
            if min_keywords_required < 1:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Invalid Minimum Score",
                                   "Minimum keyword score must be a positive integer.")
            return
        if min_keywords_required > len(keywords):
            messagebox.showwarning(
                "Minimum Score Too High",
                f"You entered {len(keywords)} distinct keyword(s), but require "
                f"{min_keywords_required}. Lower the minimum or add keywords."
            )
            return

        use_blurry = self.use_blurry.get()
        allow_partial = self.use_substring_match.get()

        filetypes = []
        if self.include_docx.get():
            filetypes.append(".docx")
        if self.include_xlsx.get():
            filetypes.append(".xlsx")
        if self.include_pptx.get():
            filetypes.append(".pptx")
        if self.include_pdf.get():
            filetypes.append(".pdf")
        if not filetypes:
            messagebox.showwarning("No File Types",
                                   "Please select at least one file type to search.")
            return

        self.result_listbox.delete(0, tk.END)
        self.context_text.delete("1.0", tk.END)
        self.results.clear()
        self.result_paths.clear()
        self.keyword_count = len(keywords)
        self.keyword_display_names = {kw.lower(): kw for kw in keywords}

        all_files = []
        for folder, _, files in os.walk(directory):
            for filename in files:
                if filename.lower().endswith(tuple(filetypes)):
                    all_files.append(os.path.join(folder, filename))

        # Preserve original progress window.
        progress_popup = Toplevel(self.root)
        progress_popup.title("Searching...")
        progress_popup.transient(self.root)
        tk.Label(progress_popup, text="Searching files...").pack(padx=10, pady=5)
        progress = ttk.Progressbar(progress_popup, orient="horizontal",
                                   length=280, mode="determinate")
        progress.pack(padx=10, pady=10)
        progress["maximum"] = max(1, len(all_files))
        progress["value"] = 0
        self.root.update_idletasks()

        matching_files = []
        try:
            for idx, full_path in enumerate(all_files):
                text = extract_text(full_path)
                matched_keywords, snippets, total_hits = find_matches(
                    text, keywords, max_mismatches, allow_partial, use_blurry
                )
                score = len(matched_keywords)

                # NEW: only include documents meeting minimum keyword score.
                if score >= min_keywords_required:
                    matching_files.append({
                        "file": full_path,
                        "score": score,
                        "matched_keywords": matched_keywords,
                        "total_hits": total_hits,
                        "snippets": snippets,
                    })

                progress["value"] = idx + 1
                self.root.update_idletasks()
        finally:
            progress_popup.destroy()

        # NEW: higher keyword scores first; more hits break ties.
        matching_files.sort(
            key=lambda item: (-item["score"], -item["total_hits"], item["file"].lower())
        )

        for item in matching_files:
            full_path = item["file"]
            self.results[full_path] = item
            self.result_paths.append(full_path)
            self.result_listbox.insert(
                tk.END,
                f"[{item['score']}/{self.keyword_count} keywords; "
                f"{item['total_hits']} hits] {full_path}"
            )

        if not matching_files:
            self.context_text.insert(
                tk.END,
                "No documents met the minimum keyword score. "
                "Try lowering the minimum or adjusting the keywords."
            )

    def on_file_select(self, event):
        selection = event.widget.curselection()
        if not selection:
            return

        # The Listbox includes score text, so use a separate path list.
        index = selection[0]
        filepath = self.result_paths[index]
        item = self.results[filepath]

        self.context_text.delete("1.0", tk.END)
        self.context_text.insert(tk.END, f"File: {filepath}\n")
        self.context_text.insert(
            tk.END, f"Keyword score: {item['score']}/{self.keyword_count}\n"
        )
        self.context_text.insert(
            tk.END, f"Matching word occurrences: {item['total_hits']}\n"
        )
        keyword_names = sorted(
            self.keyword_display_names[key] for key in item["matched_keywords"]
        )
        self.context_text.insert(
            tk.END, f"Matched keywords: {', '.join(keyword_names)}\n\n"
        )

        for snippet in item["snippets"]:
            self.context_text.insert(tk.END, f"... {snippet} ...\n\n")


# ---------------- Run the App ------------------
if __name__ == "__main__":
    root = tk.Tk()
    app = SearchApp(root)
    root.mainloop()
